"""Exportación de libros de IVA a Excel (.xlsx)."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from datos import fmt_fecha

AZUL = "1F4E78"
VERDE = "E2EFDA"
NARANJA = "FCE4D6"
GRIS = "F2F2F2"

F_TITULO = Font(size=14, bold=True, color=AZUL)
F_ENCABEZADO = Font(bold=True, color="FFFFFF")
R_ENCABEZADO = PatternFill("solid", fgColor=AZUL)
F_NEGRITA = Font(bold=True)
BORDE_SUP = Border(top=Side(style="thin"), bottom=Side(style="double"))
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)

ETIQUETAS = {
    "venta": {"plural": "Ventas", "libro": "Libro de ventas", "iva": "IVA débito", "color": VERDE},
    "compra": {"plural": "Compras", "libro": "Libro de compras", "iva": "IVA crédito", "color": NARANJA},
}


# --------------------------------------------------------------------------- helpers

def _formato_moneda(cfg):
    simbolo = cfg.get("simbolo_moneda", "$").replace('"', "")
    decimales = int(cfg.get("decimales", 0))
    base = "#,##0" + ("." + "0" * decimales if decimales else "")
    return f'"{simbolo} "{base};-"{simbolo} "{base}' if simbolo else f"{base};-{base}"


def _encabezados(ws, fila, titulos, anchos):
    for col, (titulo, ancho) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=fila, column=col, value=titulo)
        c.font, c.fill, c.alignment = F_ENCABEZADO, R_ENCABEZADO, CENTRO
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.row_dimensions[fila].height = 30


def _cabecera(ws, titulo, desde, hasta, cfg):
    ws["A1"] = titulo
    ws["A1"].font = F_TITULO
    empresa = " · ".join(x for x in (cfg.get("nombre_empresa", ""), cfg.get("id_fiscal", "")) if x)
    if empresa:
        ws["A2"] = empresa
        ws["A2"].font = F_NEGRITA
    ws["A3"] = f"Período: {fmt_fecha(desde)} al {fmt_fecha(hasta)}"
    ws["A4"] = f"Generado: {datetime.now():%d/%m/%Y %H:%M}   ·   Tasa de IVA configurada: {cfg.get('tasa_iva')}%"
    ws["A4"].font = Font(italic=True, color="666666")


def _filas_datos(ws, filas, fmt, cols_moneda, col_fecha=1):
    """Agrega filas desde la fila 2 con formato; devuelve el número de la última fila."""
    for i, valores in enumerate(filas, start=2):
        ws.append(valores)
        ws.cell(row=i, column=col_fecha).number_format = "DD/MM/YYYY"
        for col in cols_moneda:
            ws.cell(row=i, column=col).number_format = fmt
        if i % 2:
            for col in range(1, len(valores) + 1):
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=GRIS)
    return len(filas) + 1


def _fila_totales(ws, ultima, n_cols, fmt, cols_suma, cols_moneda):
    tot = ultima + 1
    ws.cell(row=tot, column=1, value="TOTAL")
    for col in cols_suma:
        letra = get_column_letter(col)
        c = ws.cell(row=tot, column=col, value=f"=SUM({letra}2:{letra}{ultima})")
        if col in cols_moneda:
            c.number_format = fmt
    for col in range(1, n_cols + 1):
        c = ws.cell(row=tot, column=col)
        c.font, c.border = F_NEGRITA, BORDE_SUP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{ultima}"
    return tot


def _hoja_detalle(wb, movimientos, fmt, con_tipo):
    ws = wb.create_sheet("Detalle")
    titulos = ["Fecha", "Tipo", "N° documento", "Descripción", "Tasa IVA %", "Neto", "IVA", "Total"]
    anchos = [12, 10, 16, 40, 11, 16, 16, 16]
    if not con_tipo:
        del titulos[1], anchos[1]
    _encabezados(ws, 1, titulos, anchos)
    desfase = 1 if con_tipo else 0
    filas = []
    for m in movimientos:
        fila = [datetime.fromisoformat(m["fecha"]).date(), m["documento"], m["descripcion"],
                m["tasa"] / 100, m["neto"], m["iva"], m["total"]]
        if con_tipo:
            fila.insert(1, m["tipo"].capitalize())
        filas.append(fila)
    ultima = _filas_datos(ws, filas, fmt, [c + desfase for c in (5, 6, 7)])
    for i, m in enumerate(movimientos, start=2):
        ws.cell(row=i, column=4 + desfase).number_format = "0.##%"
        if con_tipo:
            ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=ETIQUETAS[m["tipo"]]["color"])
    if con_tipo:
        # mezcla ventas y compras: un total sumado no tendría sentido
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{ultima}"
    else:
        _fila_totales(ws, ultima, len(titulos), fmt, (5, 6, 7), (5, 6, 7))
    return ws


# --------------------------------------------------------------------------- libro combinado

def exportar_excel(ruta, db, desde, hasta, etiqueta, cfg):
    """Libro de IVA con ventas y compras: IVA débito − IVA crédito."""
    dias = db.resumen_diario(desde, hasta)
    movimientos = db.listar(desde, hasta)
    if not dias:
        raise ValueError("No hay movimientos en el período seleccionado.")
    fmt = _formato_moneda(cfg)
    wb = Workbook()

    wd = wb.active
    wd.title = "Resumen diario"
    _encabezados(
        wd, 1,
        ["Fecha", "N° ventas", "Ventas neto", "IVA débito\n(ventas)", "Ventas total",
         "N° compras", "Compras neto", "IVA crédito\n(compras)", "Compras total", "IVA a pagar\n(débito − crédito)"],
        [12, 10, 16, 16, 16, 11, 16, 16, 16, 18],
    )
    filas = [[d["fecha"], d["n_ventas"], d["ventas_neto"], d["ventas_iva"], d["ventas_total"],
              d["n_compras"], d["compras_neto"], d["compras_iva"], d["compras_total"], f"=D{i}-H{i}"]
             for i, d in enumerate(dias, start=2)]
    moneda = (3, 4, 5, 7, 8, 9, 10)
    ultima = _filas_datos(wd, filas, fmt, moneda)
    tot = _fila_totales(wd, ultima, 10, fmt, range(2, 11), moneda)

    _hoja_detalle(wb, movimientos, fmt, con_tipo=True)

    wr = wb.create_sheet("Resumen", 0)
    _cabecera(wr, f"Libro de IVA — {etiqueta}", desde, hasta, cfg)
    _encabezados(wr, 6, ["Concepto", "Cantidad", "Neto", "IVA", "Total"], [34, 12, 18, 18, 18])
    ref = "'Resumen diario'!"
    for r, (tipo, cols) in enumerate((("venta", "BCDE"), ("compra", "FGHI")), start=7):
        e = ETIQUETAS[tipo]
        wr.cell(row=r, column=1, value=f"{e['plural']} ({e['iva']})").fill = PatternFill("solid", fgColor=e["color"])
        for col, letra in enumerate(cols, start=2):
            c = wr.cell(row=r, column=col, value=f"={ref}{letra}{tot}")
            if col > 2:
                c.number_format = fmt

    wr["A10"] = "IVA a pagar (débito − crédito)"
    wr["D10"] = "=D7-D8"
    wr["D10"].number_format = fmt
    for c in (wr["A10"], wr["D10"]):
        c.font = Font(bold=True, size=12)
        c.border = BORDE_SUP
    wr["A11"] = '=IF(D10>=0,"Resultado: IVA a pagar","Resultado: saldo a favor (remanente de crédito)")'
    wr["A11"].font = Font(italic=True, color=AZUL)
    wr["A12"] = "Días con movimientos"
    wr["B12"] = len(dias)

    wb.save(ruta)


# --------------------------------------------------------------------------- libro de ventas / compras

def exportar_libro(ruta, db, tipo, desde, hasta, etiqueta, cfg):
    """Libro exclusivo de ventas o de compras."""
    e = ETIQUETAS[tipo]
    pref = "ventas" if tipo == "venta" else "compras"
    dias = [d for d in db.resumen_diario(desde, hasta) if d[f"n_{pref}"]]
    movimientos = db.listar(desde, hasta, tipo)
    if not movimientos:
        raise ValueError(f"No hay {pref} en el período seleccionado.")
    fmt = _formato_moneda(cfg)
    wb = Workbook()

    wd = wb.active
    wd.title = "Por día"
    _encabezados(wd, 1, ["Fecha", f"N° {pref}", "Neto", e["iva"], "Total"], [12, 11, 18, 18, 18])
    filas = [[d["fecha"], d[f"n_{pref}"], d[f"{pref}_neto"], d[f"{pref}_iva"], d[f"{pref}_total"]] for d in dias]
    ultima = _filas_datos(wd, filas, fmt, (3, 4, 5))
    tot = _fila_totales(wd, ultima, 5, fmt, range(2, 6), (3, 4, 5))

    _hoja_detalle(wb, movimientos, fmt, con_tipo=False)

    wr = wb.create_sheet("Resumen", 0)
    _cabecera(wr, f"{e['libro']} — {etiqueta}", desde, hasta, cfg)
    _encabezados(wr, 6, ["Concepto", "Cantidad", "Neto", e["iva"], "Total"], [34, 12, 18, 18, 18])
    wr.cell(row=7, column=1, value=e["plural"]).fill = PatternFill("solid", fgColor=e["color"])
    for col, letra in enumerate("BCDE", start=2):
        c = wr.cell(row=7, column=col, value=f"='Por día'!{letra}{tot}")
        c.font = F_NEGRITA
        if col > 2:
            c.number_format = fmt
    wr["A9"] = "Días con movimientos"
    wr["B9"] = len(dias)

    wb.save(ruta)
